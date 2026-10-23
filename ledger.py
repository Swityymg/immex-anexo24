from __future__ import annotations

import calendar
import io
import re
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl import Workbook


DEFAULT_WORKBOOK = Path(__file__).parent / "data" / "IMMEX_Template_v2.xlsx"
ORACLES = {
    "240-3210-6000101": 23942.72,
    "240-3210-6000088": 11672.60,
    "240-3210-6000212": 900.00,
    "240-3210-6000350": 1200.00,
    "240-3210-6000571": 800.00,
}


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    if isinstance(value, str) and value.strip():
        return pd.to_datetime(value).date()
    return None


def _add_months(value: date, months: int) -> date:
    month_index = value.month - 1 + months
    year = value.year + month_index // 12
    month = month_index % 12 + 1
    day = min(value.day, calendar.monthrange(year, month)[1])
    return date(year, month, day)


def _records(sheet: Any, start_row: int, columns: tuple[int, ...]) -> list[tuple[Any, ...]]:
    result = []
    for row in sheet.iter_rows(min_row=start_row, values_only=True):
        selected = tuple(row[index - 1] if len(row) >= index else None for index in columns)
        if any(value is not None for value in selected):
            result.append(selected)
    return result


def _normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def load_csv_source_data(upload: Any) -> dict[str, Any]:
    upload.seek(0)
    frame = pd.read_csv(upload, sep=None, engine="python", encoding="utf-8-sig", dtype=str)
    aliases = {
        "pedimento": {"pedimento", "numero_pedimento", "numero_de_pedimento", "num_pedimento"},
        "clave": {"clave", "clave_pedimento", "clave_operacion"},
        "fecha_entrada": {"fecha_entrada", "fecha", "fecha_importacion", "fecha_de_importacion", "entry_date"},
        "material": {"material", "clave_material", "sku", "item"},
        "cantidad_kg": {"cantidad_kg", "cantidad", "kg", "quantity"},
        "valor_unit_usd": {"valor_unit_usd", "usd_kg", "valor_usd_kg", "valor_unitario_usd"},
        "fraccion": {"fraccion", "fraccion_arancelaria", "hts"},
        "descripcion": {"descripcion", "descripcion_material", "material_description"},
        "plazo_meses": {"plazo_meses", "plazo_legal_meses"},
        "rfc": {"rfc"},
        "razon_social": {"razon_social", "empresa", "company"},
    }
    normalized = {_normalize_header(column): column for column in frame.columns}
    mapped = {
        target: next((normalized[candidate] for candidate in candidates if candidate in normalized), None)
        for target, candidates in aliases.items()
    }
    required = ["pedimento", "fecha_entrada", "material", "cantidad_kg"]
    missing = [key for key in required if mapped[key] is None]
    if missing:
        expected = ", ".join(required)
        raise ValueError(f"No pude mapear columnas: {', '.join(missing)}. Se esperan al menos: {expected}.")

    entries = []
    for index, (_, row) in enumerate(frame.iterrows(), start=1):
        material = str(row[mapped["material"]]).strip()
        pedimento = str(row[mapped["pedimento"]]).strip()
        if not material or material.lower() == "nan" or not pedimento or pedimento.lower() == "nan":
            continue
        entries.append({
            "pedimento": pedimento,
            "lote_id": f"csv-entry-{index}",
            "clave": str(row[mapped["clave"]]).strip() if mapped["clave"] and pd.notna(row[mapped["clave"]]) else "IN",
            "fecha_entrada": pd.to_datetime(row[mapped["fecha_entrada"]]).date(),
            "material": material,
            "cantidad_kg": float(row[mapped["cantidad_kg"]]),
            "valor_unit_usd": float(row[mapped["valor_unit_usd"]]) if mapped["valor_unit_usd"] and pd.notna(row[mapped["valor_unit_usd"]]) else 0.0,
        })
    if not entries:
        raise ValueError("El CSV no contiene filas de entradas utilizables.")

    material_rows = []
    for material, group in frame.groupby(frame[mapped["material"]].astype(str).str.strip(), dropna=True):
        first = group.iloc[0]
        material_rows.append({
            "material": material,
            "descripcion": first[mapped["descripcion"]] if mapped["descripcion"] else material,
            "fraccion": str(first[mapped["fraccion"]]) if mapped["fraccion"] and pd.notna(first[mapped["fraccion"]]) else "",
            "unidad": "KGM",
            "plazo_meses": int(first[mapped["plazo_meses"]]) if mapped["plazo_meses"] and pd.notna(first[mapped["plazo_meses"]]) else 18,
        })
    rfc = "CSV-PERFIL"
    if mapped["rfc"] and frame[mapped["rfc"]].notna().any():
        rfc = str(frame[mapped["rfc"]].dropna().iloc[0]).strip()
    company = "Perfil importado por CSV"
    if mapped["razon_social"] and frame[mapped["razon_social"]].notna().any():
        company = str(frame[mapped["razon_social"]].dropna().iloc[0]).strip()
    return {
        "params": {"RFC": rfc, "Razón social": company, "Fecha de corte": date.today(),
                   "Tipo de cambio MXN por USD": 18.5, "Tolerancia de merma (alerta interna)": 0.03},
        "materials": pd.DataFrame(material_rows, columns=["material", "descripcion", "fraccion", "unidad", "plazo_meses"]),
        "bom": pd.DataFrame(columns=["producto", "material", "kg_por_pieza"]),
        "entries": pd.DataFrame(entries, columns=["pedimento", "lote_id", "clave", "fecha_entrada", "material", "cantidad_kg", "valor_unit_usd"]),
        "production": pd.DataFrame(columns=["periodo", "producto", "piezas", "merma_pct"]),
        "exits": pd.DataFrame(columns=["pedimento", "clave", "fecha", "material", "tipo", "cantidad_kg", "observacion"]),
        "profile_name": getattr(upload, "name", "perfil.csv"),
        "baseline_oracles": False,
    }


def load_custom_xlsx_source_data(upload: Any) -> dict[str, Any]:
    upload.seek(0)
    try:
        return load_source_data(upload)
    except ValueError as error:
        if "Faltan hojas requeridas" not in str(error):
            raise
    upload.seek(0)
    frame = pd.read_excel(upload, dtype=str)
    csv_buffer = io.BytesIO(frame.to_csv(index=False).encode("utf-8"))
    csv_buffer.name = getattr(upload, "name", "perfil.xlsx")
    return load_csv_source_data(csv_buffer)


def create_empty_template() -> bytes:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "LEEME"
    readme.append(["Plantilla modelo Anexo 24", "Capture solo entradas y no altere los encabezados."])

    parameters = workbook.create_sheet("Parametros")
    parameters.append(["Parámetros"])
    parameters.append([])
    parameters.append(["Parámetro", "Valor", "Tipo", "Nota"])
    for label in ["Razón social", "RFC", "Programa IMMEX (No. SE)", "Fecha de corte",
                  "Tipo de cambio MXN por USD", "Tolerancia de merma (alerta interna)"]:
        parameters.append([label, None, "CAPTURA", None])

    catalogs = workbook.create_sheet("Catalogos")
    catalogs.cell(3, 1, "Materiales")
    catalogs.cell(3, 8, "Lista de materiales (BOM)")
    for col, value in enumerate(["Clave", "Descripción", "Fracción", "UM TIGIE", "Plazo legal (meses)", "Fundamento"], start=1):
        catalogs.cell(4, col, value)
    for col, value in enumerate(["Producto", "Material", "kg por pieza"], start=8):
        catalogs.cell(4, col, value)

    entries = workbook.create_sheet("Entradas")
    for col, value in enumerate(["Pedimento", "Clave", "Fecha de entrada", "Material", "Descripción", "Fracción", "Cantidad (kg)", "Valor unitario (USD/kg)"], start=1):
        entries.cell(4, col, value)
    production = workbook.create_sheet("Consumo")
    for col, value in enumerate(["Periodo", "Producto", "Piezas producidas", "Merma real (%)"], start=1):
        production.cell(4, col, value)
    exits = workbook.create_sheet("Salidas")
    for col, value in enumerate(["Pedimento", "Clave", "Fecha", "Material", "Tipo", "Cantidad (kg)", "Observación"], start=1):
        exits.cell(4, col, value)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def load_source_data(path: Any = DEFAULT_WORKBOOK) -> dict[str, Any]:
    uploaded_stream = hasattr(path, "read")
    if uploaded_stream:
        workbook_path = None
        path.seek(0)
        workbook = load_workbook(path, data_only=True, read_only=True)
        profile_name = getattr(path, "name", "Libro cargado")
    else:
        workbook_path = Path(path)
        if not workbook_path.exists():
            raise FileNotFoundError(f"No se encontró el libro de captura: {workbook_path}")
        workbook = load_workbook(workbook_path, data_only=True, read_only=True)
        profile_name = workbook_path.name
    required_sheets = {"Parametros", "Catalogos", "Entradas", "Consumo", "Salidas"}
    missing = required_sheets.difference(workbook.sheetnames)
    if missing:
        raise ValueError(f"Faltan hojas requeridas: {', '.join(sorted(missing))}")

    params = {
        str(key).strip(): value
        for key, value in _records(workbook["Parametros"], 4, (1, 2))
        if key is not None
    }
    if "Fecha de corte" in params:
        params["Fecha de corte"] = _as_date(params["Fecha de corte"])
    materials = []
    bom = []
    for row in workbook["Catalogos"].iter_rows(min_row=5, values_only=True):
        if row[0] is not None:
            materials.append(
                {
                    "material": str(row[0]).strip(),
                    "descripcion": row[1],
                    "fraccion": str(row[2]) if row[2] is not None else "",
                    "unidad": row[3],
                    "plazo_meses": int(row[4] or 18),
                }
            )
        if len(row) > 9 and row[7] and row[8] and row[9] is not None:
            bom.append(
                {
                    "producto": str(row[7]).strip(),
                    "material": str(row[8]).strip(),
                    "kg_por_pieza": float(row[9]),
                }
            )

    entries = []
    for entry_index, (pedimento, clave, fecha, material, cantidad, valor_unitario) in enumerate(
        _records(workbook["Entradas"], 5, (1, 2, 3, 4, 7, 8)), start=1
    ):
        if not pedimento or not material:
            continue
        entries.append(
            {
                "pedimento": str(pedimento).strip(),
                "lote_id": f"template-entry-{entry_index}",
                "clave": clave,
                "fecha_entrada": _as_date(fecha),
                "material": str(material).strip(),
                "cantidad_kg": float(cantidad or 0),
                "valor_unit_usd": float(valor_unitario or 0),
            }
        )

    production = []
    production_rows = _records(workbook["Consumo"], 5, (1, 2, 3, 4))
    for periodo, producto, piezas, merma_pct in production_rows:
        if not periodo or not producto or piezas is None:
            continue
        if not isinstance(periodo, str) or re.fullmatch(r"\d{4}-\d{2}", periodo.strip()) is None:
            continue
        try:
            pieces = float(piezas)
            scrap_rate = float(merma_pct or 0)
        except (TypeError, ValueError):
            continue
        production.append(
            {
                "periodo": str(periodo).strip(),
                "producto": str(producto).strip(),
                "piezas": pieces,
                "merma_pct": scrap_rate,
            }
        )

    exits = []
    for pedimento, clave, fecha, material, tipo, cantidad, observacion in _records(
        workbook["Salidas"], 5, (1, 2, 3, 4, 5, 6, 7)
    ):
        if not pedimento or not material:
            continue
        exits.append(
            {
                "pedimento": str(pedimento).strip(),
                "clave": clave,
                "fecha": _as_date(fecha),
                "material": str(material).strip(),
                "tipo": tipo,
                "cantidad_kg": float(cantidad or 0),
                "observacion": observacion,
            }
        )

    workbook.close()
    return {
        "params": params,
        "materials": pd.DataFrame(materials, columns=["material", "descripcion", "fraccion", "unidad", "plazo_meses"]),
        "bom": pd.DataFrame(bom, columns=["producto", "material", "kg_por_pieza"]),
        "entries": pd.DataFrame(entries, columns=["pedimento", "lote_id", "clave", "fecha_entrada", "material", "cantidad_kg", "valor_unit_usd"]),
        "production": pd.DataFrame(production, columns=["periodo", "producto", "piezas", "merma_pct"]),
        "exits": pd.DataFrame(exits, columns=["pedimento", "clave", "fecha", "material", "tipo", "cantidad_kg", "observacion"]),
        "profile_name": profile_name,
        "baseline_oracles": workbook_path is not None and workbook_path.resolve() == DEFAULT_WORKBOOK.resolve(),
    }


def calculate_consumption(production: pd.DataFrame, bom: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "periodo", "producto", "material", "piezas", "kg_por_pieza",
        "consumo_teorico_kg", "merma_pct", "merma_kg", "consumo_real_kg",
    ]
    if production.empty or bom.empty:
        return pd.DataFrame(columns=columns)
    used = production.merge(bom, on="producto", how="left", validate="many_to_many")
    if used["material"].isna().any():
        unknown = sorted(used.loc[used["material"].isna(), "producto"].unique())
        raise ValueError(f"Productos sin BOM: {', '.join(unknown)}")
    used["consumo_teorico_kg"] = used["piezas"] * used["kg_por_pieza"]
    used["merma_kg"] = used["consumo_teorico_kg"] * used["merma_pct"]
    used["consumo_real_kg"] = used["consumo_teorico_kg"] + used["merma_kg"]
    return used[columns]


def calculate_ledger(
    source: dict[str, Any], fecha_corte: date, tipo_cambio: float, tolerancia_merma: float
) -> dict[str, Any]:
    entries = source["entries"].copy()
    materials = source["materials"].copy()
    consumption = calculate_consumption(source["production"], source["bom"])
    consumption_by_material = (
        consumption.groupby("material")["consumo_real_kg"].sum().to_dict()
        if not consumption.empty else {}
    )
    exits = source["exits"]
    exits_by_material = exits.groupby("material")["cantidad_kg"].sum().to_dict() if not exits.empty else {}
    required_by_material = {
        material: float(consumption_by_material.get(material, 0)) + float(exits_by_material.get(material, 0))
        for material in set(consumption_by_material) | set(exits_by_material)
    }

    if "lote_id" not in entries.columns:
        entries["lote_id"] = [f"source-entry-{index}" for index in range(len(entries))]
    allocations: dict[str, float] = {}
    shortages: dict[str, float] = {}
    for material in entries["material"].dropna().unique():
        remaining = required_by_material.get(material, 0.0)
        lots = entries.loc[entries["material"] == material].sort_values(
            ["fecha_entrada", "pedimento"], kind="stable"
        )
        for index in lots.index:
            downloaded = min(float(entries.at[index, "cantidad_kg"]), max(remaining, 0.0))
            allocations[str(entries.at[index, "lote_id"])] = downloaded
            remaining -= downloaded
        shortages[str(material)] = max(remaining, 0.0)

    catalog = materials.set_index("material").to_dict("index")
    rows = []
    for _, lot in entries.iterrows():
        material = str(lot["material"])
        item = catalog.get(material, {})
        downloaded = allocations.get(str(lot["lote_id"]), 0.0)
        balance = float(lot["cantidad_kg"]) - downloaded
        expiry = _add_months(lot["fecha_entrada"], int(item.get("plazo_meses", 18)))
        days_left = (expiry - fecha_corte).days
        if balance <= 0:
            status = "Agotado"
        elif days_left < 0:
            status = "VENCIDO"
        elif days_left <= 30:
            status = "ALERTA 30"
        elif days_left <= 60:
            status = "ALERTA 60"
        elif days_left <= 90:
            status = "ALERTA 90"
        else:
            status = "OK"
        rows.append(
            {
                "pedimento": lot["pedimento"], "lote_id": lot["lote_id"], "clave": lot["clave"],
                "fecha_entrada": lot["fecha_entrada"], "material": material,
                "descripcion": item.get("descripcion", ""), "fraccion": item.get("fraccion", ""),
                "cantidad_importada_kg": float(lot["cantidad_kg"]),
                "valor_unit_usd": float(lot["valor_unit_usd"]), "descargado_kg": downloaded,
                "saldo_kg": balance, "vencimiento": expiry, "dias_restantes": days_left,
                "estatus": status, "riesgo_mxn": balance * float(lot["valor_unit_usd"]) * tipo_cambio,
            }
        )
    ledger = pd.DataFrame(rows, columns=[
        "pedimento", "lote_id", "clave", "fecha_entrada", "material", "descripcion", "fraccion",
        "cantidad_importada_kg", "valor_unit_usd", "descargado_kg", "saldo_kg", "vencimiento",
        "dias_restantes", "estatus", "riesgo_mxn",
    ])

    names = {row["material"]: row["descripcion"] for row in materials.to_dict("records")}
    reconciliation_rows = []
    for material in sorted(required_by_material):
        imported = float(entries.loc[entries["material"] == material, "cantidad_kg"].sum())
        required = required_by_material[material]
        shortfall = shortages.get(material, max(required - imported, 0.0))
        material_exits = exits.loc[exits["material"] == material] if not exits.empty else exits
        exit_details = [
            {
                "pedimento": str(row["pedimento"]),
                "tipo": str(row["tipo"]),
                "cantidad_kg": float(row["cantidad_kg"]),
                "observacion": str(row.get("observacion", "") or ""),
            }
            for _, row in material_exits.iterrows()
        ]
        reconciliation_rows.append(
            {
                "material": material, "descripcion": names.get(material, ""),
                "consumo_real_kg": float(consumption_by_material.get(material, 0)),
                "salidas_kg": float(exits_by_material.get(material, 0)),
                "descarga_requerida_kg": required, "importado_kg": imported,
                "faltante_kg": shortfall,
                "salidas_detalle": exit_details,
                "estatus": "Faltante por conciliar" if shortfall > 1e-9 else "Conciliado",
            }
        )
    reconciliation = pd.DataFrame(reconciliation_rows, columns=[
        "material", "descripcion", "consumo_real_kg", "salidas_kg", "descarga_requerida_kg",
        "importado_kg", "faltante_kg", "salidas_detalle", "estatus",
    ])

    if consumption.empty:
        merma = pd.DataFrame(
            columns=["producto", "material", "consumo_teorico_kg", "merma_kg", "merma_pct", "tolerancia_pct", "estatus"]
        )
    else:
        merma = consumption.groupby(["producto", "material"], as_index=False).agg(
            consumo_teorico_kg=("consumo_teorico_kg", "sum"), merma_kg=("merma_kg", "sum")
        )
        merma["merma_pct"] = merma["merma_kg"] / merma["consumo_teorico_kg"]
        merma["tolerancia_pct"] = tolerancia_merma
        merma["estatus"] = merma["merma_pct"].map(
            lambda value: "MERMA EXCEDE" if value > tolerancia_merma else "Dentro de tolerancia"
        )

    alert_statuses = {"VENCIDO", "ALERTA 30", "ALERTA 60", "ALERTA 90"}
    total_mxn = float(ledger["riesgo_mxn"].sum()) if not ledger.empty else 0.0
    alert_ledger = ledger.loc[ledger["estatus"].isin(alert_statuses)]
    metrics = {
        "saldo_total_mxn": total_mxn,
        "riesgo_alertas_mxn": float(alert_ledger["riesgo_mxn"].sum()),
        "pedimentos_alerta": int(alert_ledger["pedimento"].nunique()),
        "faltante_total_kg": float(reconciliation["faltante_kg"].sum()) if not reconciliation.empty else 0.0,
    }

    checks = {}
    if source.get("baseline_oracles", False):
        actual = ledger.groupby("pedimento")["descargado_kg"].sum().to_dict()
        for pedimento, expected in ORACLES.items():
            calculated = round(float(actual.get(pedimento, 0)), 2)
            checks[pedimento] = {"esperado_kg": expected, "calculado_kg": calculated, "ok": calculated == expected}
        mp003 = round(float(reconciliation.loc[reconciliation["material"] == "MP-003", "faltante_kg"].sum()), 2)
        checks["faltante_MP-003"] = {"esperado_kg": 332.80, "calculado_kg": mp003, "ok": mp003 == 332.80}
        checks["saldo_total_mxn"] = {
            "esperado_mxn": 1592905.83, "calculado_mxn": round(metrics["saldo_total_mxn"], 2),
            "ok": round(metrics["saldo_total_mxn"], 2) == 1592905.83,
        }
        checks["riesgo_alertas_mxn"] = {
            "esperado_mxn": 425185.83, "calculado_mxn": round(metrics["riesgo_alertas_mxn"], 2),
            "ok": round(metrics["riesgo_alertas_mxn"], 2) == 425185.83,
        }
        checks["pedimentos_alerta"] = {
            "esperado": 3, "calculado": metrics["pedimentos_alerta"], "ok": metrics["pedimentos_alerta"] == 3,
        }
    return {
        "ledger": ledger, "consumption": consumption, "reconciliation": reconciliation,
        "merma": merma, "metrics": metrics, "oracle_checks": checks,
        "tolerancia_merma": tolerancia_merma, "tipo_cambio": tipo_cambio, "fecha_corte": fecha_corte,
    }